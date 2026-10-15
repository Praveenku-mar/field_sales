# Copyright (c) 2026, Your Company and contributors
# For license information, please see license.txt

import csv
import datetime
from io import BytesIO, StringIO

import frappe
from frappe.tests import IntegrationTestCase
from frappe.utils import nowdate
from openpyxl import Workbook, load_workbook

from field_sales.visit_import import (
	create_visits_from_import,
	download_visit_template,
	preview_visit_import,
)

HEADER = ["Sales Rep", "Customer", "Visit Date", "Purpose", "Notes"]


class TestVisitImport(IntegrationTestCase):
	def setUp(self):
		frappe.set_user("Administrator")
		self.manager_user = self.make_user("import.manager@example.com", ["Field Sales Rep", "Sales Manager"])
		self.manager_rep = self.make_rep(self.manager_user, "Import Manager")
		self.rep_user = self.make_user("import.rep@example.com", ["Field Sales Rep"])
		self.rep = self.make_rep(self.rep_user, "Import Rep", reports_to=self.manager_rep.name)
		self.other_user = self.make_user("import.other@example.com", ["Field Sales Rep"])
		self.other_rep = self.make_rep(self.other_user, "Import Other")
		self.shop = self.make_retailer("Import Test Stores", "Coimbatore")
		self.twin_a = self.make_retailer("Twin Traders", "Chennai")
		self.twin_b = self.make_retailer("Twin Traders", "Madurai")
		self.files = 0

	def tearDown(self):
		frappe.set_user("Administrator")
		# IntegrationTestCase only rolls back once per class.
		frappe.db.after_rollback.reset()
		frappe.db.rollback()

	def make_user(self, email, roles):
		user = frappe.get_doc(
			{"doctype": "User", "email": email, "first_name": email.split("@")[0], "send_welcome_email": 0}
		)
		user.insert(ignore_permissions=True)
		user.add_roles(*roles)
		return user.name

	def make_rep(self, user, full_name, reports_to=None):
		return frappe.get_doc(
			{"doctype": "Sales Rep", "user": user, "full_name": full_name, "reports_to": reports_to}
		).insert(ignore_permissions=True)

	def make_retailer(self, name, city):
		return frappe.get_doc(
			{
				"doctype": "Retailer",
				"retailer_name": name,
				"phone": "9876543210",
				"address_line1": "1 Main Road",
				"city": city,
				"sales_rep": self.rep.name,
			}
		).insert(ignore_permissions=True)

	def make_visit(self, rep, retailer, date, status="Planned"):
		return frappe.get_doc(
			{"doctype": "Visit", "sales_rep": rep, "retailer": retailer, "visit_date": date, "status": status}
		).insert(ignore_permissions=True)

	def upload_xlsx(self, rows, user="Administrator"):
		wb = Workbook()
		ws = wb.active
		ws.title = "Visits"
		for row in [HEADER, *rows]:
			ws.append(row)
		buf = BytesIO()
		wb.save(buf)
		return self.save_file("visits.xlsx", buf.getvalue(), user)

	def upload_csv(self, rows, user="Administrator"):
		buf = StringIO()
		csv.writer(buf).writerows([HEADER, *rows])
		return self.save_file("visits.csv", buf.getvalue().encode(), user)

	def save_file(self, filename, content, user):
		self.files += 1
		file = frappe.get_doc(
			{
				"doctype": "File",
				"file_name": f"{self.files}-{filename}",
				"content": content,
				"is_private": 1,
				"owner": user,
			}
		).insert(ignore_permissions=True)
		file.db_set("owner", user)
		return file.file_url

	def preview_as(self, user, url):
		frappe.set_user(user)
		try:
			return preview_visit_import(url)
		finally:
			frappe.set_user("Administrator")

	def test_template_has_both_sheets(self):
		download_visit_template()
		wb = load_workbook(BytesIO(frappe.response["filecontent"]))
		self.assertEqual(wb.sheetnames, ["Visits", "Lists"])
		self.assertEqual([c.value for c in wb["Visits"][1]], HEADER)
		lists_header = [c.value for c in wb["Lists"][1]]
		self.assertEqual(
			lists_header, ["Sales Rep email", "Sales Rep name", "Customer ID", "Customer name", "City", "Purposes"]
		)
		self.assertIn(self.rep.name, [r[0] for r in wb["Lists"].iter_rows(min_row=2, values_only=True)])

	def test_template_limited_to_hierarchy_for_reps(self):
		frappe.set_user(self.manager_user)
		try:
			download_visit_template()
		finally:
			frappe.set_user("Administrator")
		wb = load_workbook(BytesIO(frappe.response["filecontent"]))
		emails = {r[0] for r in wb["Lists"].iter_rows(min_row=2, values_only=True)}
		self.assertIn(self.rep.name, emails)
		self.assertNotIn(self.other_rep.name, emails)

	def test_preview_resolves_by_id_and_name(self):
		url = self.upload_xlsx(
			[
				[self.rep.name, self.shop.name, datetime.datetime(2030, 1, 5), "collection", "Bring invoice"],
				[self.rep.full_name.upper(), "import test stores", "06-01-2030", None, None],
				[self.rep.name, self.shop.name, "2030/01/07", "", ""],
			]
		)
		result = preview_visit_import(url)
		self.assertEqual((result["ready"], result["errors"]), (2, 1))
		first, second, third = result["rows"]
		self.assertEqual(first["visit_date"], "2030-01-05")
		self.assertEqual(first["visit_purpose"], "Collection")
		self.assertEqual(first["retailer"], self.shop.name)
		self.assertEqual(first["notes"], "Bring invoice")
		self.assertEqual((second["sales_rep"], second["retailer"], second["visit_purpose"]), (self.rep.name, self.shop.name, "Order Taking"))
		self.assertEqual(second["visit_date"], "2030-01-06")
		self.assertIn("not a date", third["error"])

	def test_preview_flags_bad_rows(self):
		url = self.upload_xlsx(
			[
				["nobody", self.shop.name, "05-01-2030", "", ""],
				[self.rep.name, "No Such Shop", "05-01-2030", "", ""],
				[self.rep.name, "Twin Traders", "05-01-2030", "", ""],
				[self.rep.name, self.shop.name, "05-01-2030", "Picnic", ""],
				[self.rep.name, self.shop.name, "", "", ""],
			]
		)
		rows = preview_visit_import(url)["rows"]
		self.assertIn("not found", rows[0]["error"])
		self.assertIn("not found", rows[1]["error"])
		self.assertIn("Customer ID", rows[2]["error"])
		self.assertIn("Chennai", rows[2]["error"])
		self.assertIn("Madurai", rows[2]["error"])
		self.assertIn("Purpose 'Picnic'", rows[3]["error"])
		self.assertIn("not a date", rows[4]["error"])
		# An ambiguous name is fine once the ID is used.
		ok = preview_visit_import(self.upload_xlsx([[self.rep.name, self.twin_a.name, "05-01-2030", "", ""]]))
		self.assertEqual(ok["ready"], 1)

	def test_preview_flags_duplicates(self):
		existing = self.make_visit(self.rep.name, self.shop.name, "2030-02-01")
		cancelled = self.make_visit(self.rep.name, self.shop.name, "2030-02-02", status="Cancelled")
		url = self.upload_xlsx(
			[
				[self.rep.name, self.shop.name, "01-02-2030", "", ""],
				[self.rep.name, self.shop.name, "02-02-2030", "", ""],
				[self.rep.name, self.shop.name, "03-02-2030", "", ""],
				[self.rep.name, self.shop.name, "03-02-2030", "", ""],
			]
		)
		rows = preview_visit_import(url)["rows"]
		self.assertIn(existing.name, rows[0]["error"])
		self.assertIsNone(rows[1]["error"], cancelled.name)
		self.assertIsNone(rows[2]["error"])
		self.assertIn("Duplicate of row 4", rows[3]["error"])

	def test_rep_cannot_import_for_rep_outside_hierarchy(self):
		url = self.upload_xlsx(
			[
				[self.other_rep.name, self.shop.name, "05-01-2030", "", ""],
				[self.rep.name, self.shop.name, "05-01-2030", "", ""],
			],
			user=self.manager_user,
		)
		result = self.preview_as(self.manager_user, url)
		self.assertIn("not found in your team", result["rows"][0]["error"])
		self.assertIsNone(result["rows"][1]["error"])

		frappe.set_user(self.manager_user)
		try:
			created = create_visits_from_import(url)
		finally:
			frappe.set_user("Administrator")
		self.assertEqual(len(created["created"]), 1)
		self.assertFalse(frappe.db.exists("Visit", {"sales_rep": self.other_rep.name}))

	def test_create_inserts_valid_rows_and_second_upload_is_all_duplicates(self):
		rows = [
			[self.rep.name, self.shop.name, "05-01-2030", "Complaint", "Late delivery"],
			["nobody", self.shop.name, "05-01-2030", "", ""],
			[self.rep.name, self.twin_a.name, "05-01-2030", "", ""],
		]
		result = create_visits_from_import(self.upload_xlsx(rows))
		self.assertEqual([c["row"] for c in result["created"]], [2, 4])
		self.assertEqual([s["row"] for s in result["skipped"]], [3])
		visit = frappe.get_doc("Visit", result["created"][0]["name"])
		self.assertEqual(
			(visit.status, visit.visit_purpose, visit.notes, visit.retailer_name),
			("Planned", "Complaint", "Late delivery", "Import Test Stores"),
		)
		self.assertEqual(str(visit.visit_date), "2030-01-05")

		again = preview_visit_import(self.upload_xlsx(rows[:1] + rows[2:]))
		self.assertEqual((again["ready"], again["errors"]), (0, 2))
		self.assertEqual(create_visits_from_import(self.upload_xlsx(rows[:1]))["created"], [])

	def test_create_deletes_upload_only_when_visits_were_created(self):
		url = self.upload_xlsx([["nobody", self.shop.name, "05-01-2030", "", ""]])
		create_visits_from_import(url)
		self.assertTrue(frappe.db.exists("File", {"file_url": url}))
		url = self.upload_xlsx([[self.rep.name, self.shop.name, "05-01-2030", "", ""]])
		create_visits_from_import(url)
		self.assertFalse(frappe.db.exists("File", {"file_url": url}))

	def test_csv_works(self):
		url = self.upload_csv([[self.rep.name, self.shop.name, "05/01/2030", "Follow-up", "csv note"]])
		result = preview_visit_import(url)
		self.assertEqual((result["ready"], result["errors"]), (1, 0))
		self.assertEqual(result["rows"][0]["visit_purpose"], "Follow-up")
		self.assertEqual(len(create_visits_from_import(url)["created"]), 1)

	def test_plain_rep_cannot_import(self):
		url = self.upload_xlsx([[self.rep.name, self.shop.name, "05-01-2030", "", ""]], user=self.rep_user)
		frappe.set_user(self.rep_user)
		try:
			for call in (download_visit_template, lambda: preview_visit_import(url), lambda: create_visits_from_import(url)):
				with self.assertRaises(frappe.PermissionError):
					call()
		finally:
			frappe.set_user("Administrator")
		self.assertFalse(frappe.db.exists("Visit", {"sales_rep": self.rep.name, "visit_date": "2030-01-05"}))

	def test_rejects_other_users_file_and_bad_headers(self):
		url = self.upload_xlsx([[self.rep.name, self.shop.name, "05-01-2030", "", ""]], user=self.other_user)
		with self.assertRaises(frappe.ValidationError):
			preview_visit_import(url)

		wb = Workbook()
		wb.active.append(["Foo", "Bar"])
		buf = BytesIO()
		wb.save(buf)
		with self.assertRaises(frappe.ValidationError):
			preview_visit_import(self.save_file("bad.xlsx", buf.getvalue(), "Administrator"))
